"""Build the assignment report (Word .docx) from the repo's figures and screenshots.

Usage:   python scripts/build_report.py
Output:  report/MLOps_Assignment1_Report.docx

Edit the constants and the text blocks below so the report reads in your own words.
Screenshots are embedded automatically from screenshots/ (main body + appendix).
"""
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from matplotlib.figure import Figure
from matplotlib.patches import FancyBboxPatch

# ---- Edit these ---------------------------------------------------------------------------
AUTHOR = "Your Name"
STUDENT_ID = "Your BITS ID"
REPO_URL = "https://github.com/BharathKumarABits/MLOPs"

# REPLACE these with your own observations from notebooks/01_eda.ipynb.
EDA_OBSERVATIONS = [
    "The classes are almost balanced (164 healthy, 139 with disease), so no resampling was "
    "needed and accuracy is a reasonable metric next to recall and ROC-AUC.",
    "thal (0.53), ca (0.46), exang (0.43), oldpeak (0.43) and cp (0.41) correlate most "
    "strongly with the target; thalach correlates negatively (-0.42), meaning patients "
    "who reach a higher maximum heart rate are less likely to have the disease.",
    "fbs (0.03) and chol (0.09) show almost no linear relationship with the target on their "
    "own, but they were kept: a weak feature can still help a multivariate model.",
    "cp, restecg, slope and thal are category codes, so Pearson correlations for them are "
    "only a rough guide; this is why they are one-hot encoded in the pipeline.",
    "Only 6 values are missing (4 in ca, 2 in thal, about 2% of rows), so imputation inside "
    "the pipeline is a safe choice.",
]
# -------------------------------------------------------------------------------------------

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "screenshots"
FIGS = ROOT / "report" / "figures"
OUT = ROOT / "report" / "MLOps_Assignment1_Report.docx"
NAVY = RGBColor(0x1F, 0x2A, 0x6B)

FEATURE_TABLE = [
    ("age", "Age in years", "numeric"),
    ("sex", "1 = male, 0 = female", "binary"),
    ("cp", "Chest pain type (1-4)", "category"),
    ("trestbps", "Resting blood pressure (mm Hg)", "numeric"),
    ("chol", "Serum cholesterol (mg/dl)", "numeric"),
    ("fbs", "Fasting blood sugar > 120 mg/dl", "binary"),
    ("restecg", "Resting ECG result (0-2)", "category"),
    ("thalach", "Maximum heart rate achieved", "numeric"),
    ("exang", "Exercise-induced angina", "binary"),
    ("oldpeak", "ST depression induced by exercise", "numeric"),
    ("slope", "Slope of the peak exercise ST segment (1-3)", "category"),
    ("ca", "Major vessels coloured by fluoroscopy (0-3)", "numeric (count)"),
    ("thal", "Thalassemia: 3 normal, 6 fixed, 7 reversible defect", "category"),
]

RESULTS_TABLE = [
    ("logreg_baseline", "0.907", "0.766", "0.958", "0.929", "0.869"),
    ("logreg_tuned (selected)", "0.908", "0.828", "0.958", "0.929", "0.869"),
    ("rf_baseline", "0.894", "0.783", "0.945", "0.964", "0.869"),
    ("rf_tuned", "0.902", "0.764", "0.955", "0.929", "0.885"),
]

# (file, caption, width in inches) for the main body, grouped by section key.
MAIN_FIGURES = {
    "mlflow": [
        ("mlflow_01_experiment_runs.png", "MLflow experiment with all four runs", 6.0),
    ],
    "ci": [
        ("ci_02_job_steps.png", "GitHub Actions job: lint, tests, training, quality gate", 5.4),
        ("ci_04_failed_run.png", "Deliberately broken test makes the pipeline fail", 5.4),
    ],
    "docker": [
        ("docker_02_run_predict.png", "Container health check and a prediction", 5.4),
        ("docker_04_ps_healthy.png", "Container reported healthy by Docker", 5.4),
    ],
    "k8s": [
        ("k8s_01_nodes_pods_svc.png", "Nodes, deployment, pods and LoadBalancer service", 6.0),
        ("k8s_02_loadbalancer_curl.png", "Requests served through the LoadBalancer address", 6.0),
    ],
    "monitoring": [
        ("mon_03_grafana_dashboard.png", "Grafana dashboard under generated traffic", 6.0),
        ("mon_01_prometheus_targets.png", "Prometheus scraping both API pods (2/2 up)", 5.0),
    ],
}


def _shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def draw_architecture(path):
    """Draw the pipeline architecture as a PNG (training lane, then serving lane)."""
    width, height = 1.8, 0.95
    fig = Figure(figsize=(11, 6.0), dpi=150)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 6.0)
    ax.axis("off")

    def box(x, y, text, color):
        ax.add_patch(FancyBboxPatch(
            (x, y), width, height, boxstyle="round,pad=0.03,rounding_size=0.08",
            facecolor=color, edgecolor="#333333", linewidth=1.2))
        ax.text(x + width / 2, y + height / 2, text, ha="center", va="center", fontsize=8)

    def arrow(start, end, style="-"):
        ax.annotate("", xy=end, xytext=start, arrowprops=dict(
            arrowstyle="->", lw=1.4, color="#333333", linestyle=style))

    xs = [0.2, 2.4, 4.6, 6.8, 9.0]
    lane_a, lane_b = 3.2, 0.8
    train = [
        ("UCI Heart Disease\ndataset (303 rows)", "#E8F0FE"),
        ("data.py\ndownload + clean", "#E8F0FE"),
        ("features.py\nimpute, scale, encode", "#E8F0FE"),
        ("train.py\n5-fold CV + tuning", "#E8F0FE"),
        ("models/\nheart_risk_model.joblib", "#FFF4CC"),
    ]
    serve = [
        ("FastAPI service\n/predict /health /metrics", "#E3F6E8"),
        ("Docker image\npython:3.11-slim", "#E3F6E8"),
        ("Minikube: 2 pods +\nLoadBalancer service", "#E3F6E8"),
        ("Prometheus\n(scrapes each pod)", "#FDE7E7"),
        ("Grafana\ndashboard", "#FDE7E7"),
    ]
    for x, (text, color) in zip(xs, train):
        box(x, lane_a, text, color)
    for x, (text, color) in zip(reversed(xs), serve):
        box(x, lane_b, text, color)
    for left, right in zip(xs, xs[1:]):
        arrow((left + width + 0.03, lane_a + height / 2), (right - 0.03, lane_a + height / 2))
        arrow((right - 0.03, lane_b + height / 2), (left + width + 0.03, lane_b + height / 2))
    arrow((xs[-1] + width / 2, lane_a - 0.03), (xs[-1] + width / 2, lane_b + height + 0.03))

    box(xs[3], 4.9, "MLflow tracking\nSQLite + artifacts", "#EFE6FB")
    arrow((xs[3] + width / 2, lane_a + height + 0.03), (xs[3] + width / 2, 4.9 - 0.03))
    box(xs[0], 4.9, "GitHub Actions CI\nlint, test, train, gate", "#EFE6FB")
    arrow((xs[0] + width / 2, 4.9 - 0.03), (xs[0] + width / 2, lane_a + height + 0.03), "--")
    ax.text(0.2, 0.25, "Training lane (top) feeds the saved model into the serving lane (bottom).",
            fontsize=8, style="italic")
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path)


class Report:
    def __init__(self):
        self.doc = Document()
        self.fig_no = 0
        self.used = set()
        self.missing = []
        self._setup()

    def _setup(self):
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.left_margin = section.right_margin = Inches(0.9)
        section.top_margin = section.bottom_margin = Inches(0.8)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(6)
        for name, size in (("Heading 1", 16), ("Heading 2", 13)):
            style = self.doc.styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.color.rgb = NAVY
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Heart Disease Risk MLOps Pipeline  |  Page ")
        for kind in ("begin", "instr", "end"):
            run = footer.add_run()
            if kind == "instr":
                element = OxmlElement("w:instrText")
                element.set(qn("xml:space"), "preserve")
                element.text = "PAGE"
            else:
                element = OxmlElement("w:fldChar")
                element.set(qn("w:fldCharType"), kind)
            run._r.append(element)

    # ---- text helpers
    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def p(self, text, italic=False):
        par = self.doc.add_paragraph()
        run = par.add_run(text)
        run.italic = italic
        return par

    def bullets(self, items):
        for item in items:
            self.doc.add_paragraph(item, style="List Bullet")

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            par = self.doc.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.left_indent = Inches(0.2)
            run = par.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def table(self, header, rows, widths=None, bold_rows=()):
        table = self.doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, text in zip(table.rows[0].cells, header):
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.bold = True
            run.font.size = Pt(9)
            _shade(cell, "DCE3F5")
        for index, row in enumerate(rows):
            cells = table.add_row().cells
            for cell, text in zip(cells, row):
                cell.text = ""
                run = cell.paragraphs[0].add_run(str(text))
                run.font.size = Pt(9)
                run.bold = index in bold_rows
        if widths:
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Inches(width)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---- figures
    def _picture(self, container, path, width):
        path = Path(path)
        self.used.add(path.name)
        par = container.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not path.exists():
            self.missing.append(path.name)
            run = par.add_run(f"[missing image: {path.name}]")
            run.italic = True
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            return
        par.add_run().add_picture(str(path), width=Inches(width))

    def _caption(self, container, text):
        self.fig_no += 1
        par = container.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run(f"Figure {self.fig_no}: {text}")
        run.italic = True
        run.font.size = Pt(9)

    def figure(self, path, caption, width=6.0):
        self._picture(self.doc, path, width)
        self._caption(self.doc, caption)

    def pair(self, first, second, width=3.1):
        table = self.doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, (path, caption) in zip(table.rows[0].cells, (first, second)):
            self._picture(cell, path, width)
            self._caption(cell, caption)

    def main_figures(self, key):
        for name, caption, width in MAIN_FIGURES[key]:
            self.figure(SHOTS / name, caption, width)

    def appendix(self):
        self.doc.add_page_break()
        self.h1("Appendix A: Additional evidence")
        self.p("Full-resolution images are in the screenshots/ folder of the repository.",
               italic=True)
        remaining = sorted(
            f for f in SHOTS.glob("*") if f.suffix.lower() in {".png", ".jpg", ".jpeg"}
            and f.name not in self.used)
        for i in range(0, len(remaining), 2):
            chunk = remaining[i:i + 2]
            items = [(f, f.stem.replace("_", " ")) for f in chunk]
            if len(items) == 1:
                self.figure(items[0][0], items[0][1], 3.1)
            else:
                self.pair(items[0], items[1])


def load_card():
    card_path = ROOT / "models" / "model_card.json"
    if card_path.exists():
        return json.loads(card_path.read_text())
    return {"selected_model": "logreg_tuned", "cv_mean": {}, "test": {}, "versions": {}}


def build():
    report = Report()
    doc = report.doc
    card = load_card()
    versions = card.get("versions", {})
    draw_architecture(FIGS / "architecture.png")

    # ---- title page
    for _ in range(5):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Heart Disease Risk Prediction")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = NAVY
    for text, size in (
        ("End-to-End MLOps Pipeline: Experiment Tracking, CI/CD, Containers, "
         "Kubernetes and Monitoring", 13),
        ("Machine Learning Operations (MLOps), AIMLCZG523, Assignment 01", 12),
        ("BITS Pilani WILP, M.Tech AI/ML", 12),
        ("", 12),
        (f"{AUTHOR}  |  {STUDENT_ID}", 12),
        (date.today().strftime("%d %B %Y"), 11),
        (f"Code repository: {REPO_URL}", 11),
    ):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run(text).font.size = Pt(size)
    doc.add_page_break()

    report.h1("Contents")
    report.bullets([
        "1. Introduction", "2. Architecture", "3. Setup and installation",
        "4. Data acquisition and exploratory analysis", "5. Feature engineering and models",
        "6. Experiment tracking with MLflow", "7. Model packaging and reproducibility",
        "8. CI/CD and automated testing", "9. Containerization", "10. Kubernetes deployment",
        "11. Monitoring and logging", "12. Challenges and lessons learned",
        "13. Limitations and future work", "14. Conclusion", "References",
        "Appendix A: Additional evidence"])

    # ---- 1. Introduction
    report.h1("1. Introduction")
    report.p(
        "The goal of this assignment is to design, build and deploy a reproducible machine "
        "learning solution using MLOps practices: automated data handling, experiment "
        "tracking, a tested CI/CD pipeline, a containerised model service, deployment on "
        "Kubernetes, and monitoring. The use case is a binary classifier that predicts the "
        "risk of heart disease from patient health data and exposes the prediction through "
        "a REST API.")
    report.p(
        "The data is the Heart Disease dataset from the UCI Machine Learning Repository "
        "(Cleveland cohort): 303 patients, 13 clinical features and a diagnosis target. "
        "The original target (0-4) is binarised to 0 = no disease and 1 = disease. This is an "
        "educational project, not a medical device.")
    report.table(
        ["Task", "Delivered in this project"],
        [
            ("1. Data and EDA", "Download script, cleaning, EDA notebook and figures"),
            ("2. Features and models", "Preprocessing pipeline, Logistic Regression and "
                                       "Random Forest, cross-validation, tuning"),
            ("3. Experiment tracking", "MLflow runs with parameters, metrics, plots, models"),
            ("4. Packaging", "Saved pipeline, model card, pinned requirements"),
            ("5. CI/CD and tests", "16 pytest tests, GitHub Actions with quality gate"),
            ("6. Containerization", "FastAPI service in a Docker image with health check"),
            ("7. Deployment", "Kubernetes manifests, 2 replicas, LoadBalancer service"),
            ("8. Monitoring", "Request logging, Prometheus metrics, Grafana dashboard"),
        ], widths=[1.8, 4.6])

    # ---- 2. Architecture
    report.h1("2. Architecture")
    report.figure(FIGS / "architecture.png", "Overall architecture of the pipeline", 6.4)
    report.p(
        "Training runs on code from the repository: the data module downloads and cleans the "
        "dataset, the feature pipeline imputes, scales and encodes it, and the training "
        "script cross-validates and tunes candidate models while logging every run to MLflow. "
        "The best pipeline is saved as one artifact. The serving lane loads that artifact in "
        "a FastAPI service, packages it as a Docker image, runs it as two Kubernetes pods "
        "behind a LoadBalancer service, and exposes metrics that Prometheus scrapes from each "
        "pod and Grafana visualises. GitHub Actions lints, tests and retrains on every push.")

    # ---- 3. Setup
    report.h1("3. Setup and installation")
    report.p(
        "Development and deployment were done on a cloud lab virtual machine provided by "
        "the institute (Rocky Linux 9.5, 4 vCPU, 15 GB RAM, Docker, Minikube) using VS Code. "
        "Python 3.11 is used everywhere: the local environment, the Docker image and CI. "
        "From a clean machine:")
    report.code(
        "git clone " + REPO_URL + ".git && cd MLOPs\n"
        "python3.11 -m venv .venv && source .venv/bin/activate\n"
        "pip install -r requirements-dev.txt\n"
        "python -m heart_risk.data     # download from UCI and write data/raw, data/processed\n"
        "python -m heart_risk.train    # train, track in MLflow, save models/\n"
        "pytest -v && flake8 heart_risk test scripts api")
    report.p("Container and cluster (the lab VM is reset between sessions, so two scripts "
             "rebuild everything):")
    report.code(
        "bash scripts/setup_vm.sh          # kubectl, GitHub CLI, VS Code extensions\n"
        "bash scripts/deploy_local.sh      # Minikube, image build, Deployment, Service\n"
        "minikube tunnel                   # second terminal; provides the LoadBalancer IP\n"
        "bash scripts/deploy_monitoring.sh # Prometheus and Grafana")

    # ---- 4. Data and EDA
    report.h1("4. Data acquisition and exploratory analysis")
    report.h2("4.1 Acquisition and cleaning")
    report.p(
        "heart_risk/data.py downloads the dataset programmatically (ucimlrepo, dataset id 45), "
        "saves the raw CSV, binarises the target, drops duplicate rows and writes the cleaned "
        "CSV. The result is 303 rows by 14 columns. Six values are missing (4 in ca, 2 in "
        "thal). They are deliberately kept in the cleaned file and imputed inside the model "
        "pipeline, so imputation statistics come from training data only and cannot leak "
        "information from validation or test data.")
    report.table(["Feature", "Meaning", "Type"], FEATURE_TABLE, widths=[0.9, 4.0, 1.5])
    report.h2("4.2 Exploratory analysis")
    report.figure(FIGS / "class_balance.png", "Class balance (164 healthy, 139 with disease)",
                  3.4)
    report.figure(FIGS / "numeric_histograms.png", "Distributions of numeric features by class",
                  6.2)
    report.figure(FIGS / "correlation_heatmap.png", "Correlation heatmap", 4.6)
    report.p("Observations:")
    report.bullets(EDA_OBSERVATIONS)

    # ---- 5. Features and models
    report.h1("5. Feature engineering and model development")
    report.h2("5.1 Preprocessing pipeline")
    report.p(
        "All preprocessing lives in one scikit-learn ColumnTransformer (heart_risk/features.py) "
        "that is fitted on training data only and stored inside the saved model. The result "
        "has 22 columns.")
    report.table(
        ["Feature group", "Columns", "Treatment"],
        [
            ("Numeric", "age, trestbps, chol, thalach, oldpeak, ca",
             "Median imputation, standard scaling"),
            ("Category codes", "cp, restecg, slope, thal",
             "Most-frequent imputation, one-hot encoding (unseen categories ignored)"),
            ("Binary", "sex, fbs, exang", "Passed through unchanged"),
        ], widths=[1.3, 2.5, 2.6])
    report.h2("5.2 Models, tuning and evaluation protocol")
    report.p(
        "The data is split 80/20 with stratification (242 training and 61 test patients, "
        "fixed seed). Model selection and tuning use stratified 5-fold cross-validation on "
        "the training part only; the test part is used once, for the final estimate. Four "
        "candidates were trained: untuned and tuned versions of Logistic Regression and "
        "Random Forest. Tuning used grid search with ROC-AUC as the refit metric while also "
        "recording accuracy, precision, recall and F1.")
    report.table(
        ["Model", "Search grid"],
        [
            ("Logistic Regression", "C in {0.01, 0.1, 1, 10, 100}; class_weight in "
                                    "{None, balanced}"),
            ("Random Forest", "n_estimators in {100, 300}; max_depth in {None, 4, 8}; "
                              "min_samples_leaf in {1, 3, 5}"),
        ], widths=[1.6, 4.8])
    report.h2("5.3 Results")
    report.table(
        ["Run", "CV ROC-AUC", "CV recall", "Test ROC-AUC", "Test recall", "Test accuracy"],
        RESULTS_TABLE, widths=[1.9, 0.9, 0.8, 1.0, 0.9, 0.9], bold_rows=(1,))
    cv = card.get("cv_mean", {})
    report.p(
        "The selected model is the tuned Logistic Regression (C = 1, class_weight = balanced), "
        "chosen by the pre-declared rule of highest mean cross-validated ROC-AUC. Its "
        f"cross-validated scores are ROC-AUC {cv.get('roc_auc', 0.908)}, recall "
        f"{cv.get('recall', 0.828)} and precision {cv.get('precision', 0.870)}. Balanced class "
        "weights raised cross-validated recall from 0.766 to 0.828, which matters in a "
        "screening setting where a missed patient is costlier than a false alarm. All four "
        "models are within the noise of each other (CV ROC-AUC 0.894 to 0.908). The test set "
        "holds only 61 patients, so a difference of one patient moves accuracy by 1.6 "
        "points; for that reason cross-validation is treated as the main estimate and the "
        "test results as confirmation.")

    # ---- 6. Experiment tracking
    report.h1("6. Experiment tracking with MLflow")
    report.p(
        "Every candidate is one MLflow run in the experiment heart-disease-risk, stored in a "
        "SQLite backend (MLflow 3 no longer maintains the file store). Each run logs the "
        "model family and tuning flag as tags, the cross-validation setup and best "
        "hyperparameters as parameters, cross-validated mean and standard deviation of five "
        "metrics plus test metrics, a ROC curve, a confusion matrix, the full tuning grid "
        "(cv_results.csv) and the fitted model with its signature and an input example.")
    report.main_figures("mlflow")
    report.pair((SHOTS / "mlflow_04_roc_curve.png", "ROC curve of the selected model"),
                (SHOTS / "mlflow_05_confusion_matrix.png", "Confusion matrix on the test split"))

    # ---- 7. Packaging
    report.h1("7. Model packaging and reproducibility")
    report.bullets([
        "Saved artifact: models/heart_risk_model.joblib is the complete pipeline "
        "(preprocessing + classifier); the same model is logged to MLflow with a signature.",
        "Model card: models/model_card.json records the selection rule, cross-validated and "
        "test metrics, the feature list and library versions (Python "
        f"{versions.get('python', '3.11')}, scikit-learn {versions.get('scikit-learn', '')}, "
        f"MLflow {versions.get('mlflow', '')}).",
        "Pinned dependencies: requirements.txt (pipeline), requirements-dev.txt (tests, lint) "
        "and requirements-api.txt (slim runtime for the container).",
        "Determinism: fixed random seeds and a stratified split; a saved model can only be "
        "loaded reliably with the same scikit-learn version, hence the exact pins.",
        "Clean-environment proof: the CI job installs the pinned requirements on a fresh "
        "runner and runs every step from scratch.",
    ])

    # ---- 8. CI/CD
    report.h1("8. CI/CD and automated testing")
    report.p(
        "16 pytest tests cover the data cleaning, the feature pipeline, the prediction "
        "module and the API (validation errors, missing optional values, metrics endpoint). "
        "Tests build their own synthetic data and a small model, so they are fast and need "
        "no network access.")
    report.table(
        ["Test file", "What it checks"],
        [
            ("test_data.py", "Target binarisation, columns, duplicate removal, missing values"),
            ("test_features.py", "22 output columns, no NaNs, unseen categories, scaling"),
            ("test_predict.py", "Output fields, confidence range, missing features rejected"),
            ("test_api.py", "/health, /predict, validation (HTTP 422), /metrics counters"),
        ], widths=[1.5, 4.9])
    report.p(
        "Statement coverage is 51% overall: 100% for the feature and prediction modules, 86% "
        "for the API, 48% for the data module (the network download is not unit tested) and "
        "0% for the training script, which is exercised end-to-end by the CI pipeline "
        "instead.")
    report.p(
        "The GitHub Actions workflow (.github/workflows/ci.yml) runs on every push: install "
        "pinned dependencies, lint with flake8, run pytest with coverage, train the model, "
        "check a quality gate (test ROC-AUC must be at least 0.80) and upload the model and "
        "reports as artifacts. Any failing step stops the pipeline with its log; this was "
        "demonstrated by deliberately breaking a test.")
    report.main_figures("ci")

    # ---- 9. Containerization
    report.h1("9. Containerization")
    report.p(
        "The model is served by a FastAPI application (api/main.py) packaged with a "
        "Dockerfile based on python:3.11-slim. It installs only the runtime dependencies, "
        "runs as a non-root user and declares a health check.")
    report.table(
        ["Endpoint", "Purpose"],
        [
            ("POST /predict", "JSON patient record in; prediction, label, confidence and "
                              "disease probability out; invalid input returns HTTP 422"),
            ("GET /health", "Liveness and model-loaded status (used by Docker and Kubernetes)"),
            ("GET /metrics", "Prometheus metrics"),
            ("GET /docs", "Interactive OpenAPI documentation"),
        ], widths=[1.4, 5.0])
    report.code(
        'curl -X POST http://localhost:8000/predict -H "Content-Type: application/json" \\\n'
        "     -d @api/sample_request.json\n"
        '{"prediction":1,"label":"heart disease","confidence":0.9812,'
        '"probability_disease":0.9812}')
    report.main_figures("docker")

    # ---- 10. Kubernetes
    report.h1("10. Kubernetes deployment")
    report.p(
        "The image is deployed to a local Minikube cluster (Kubernetes 1.32). The Deployment "
        "runs two replicas with readiness and liveness probes on /health and CPU and memory "
        "requests and limits. A Service of type LoadBalancer exposes the API on port 8000 "
        "(Minikube tunnel provides the external address) and a headless Service lets "
        "Prometheus discover every pod. The whole deployment is reproducible with "
        "scripts/deploy_local.sh.")
    report.main_figures("k8s")

    # ---- 11. Monitoring
    report.h1("11. Monitoring and logging")
    report.p(
        "The API writes one structured log line per request (request id, method, path, "
        "status, duration) and one per prediction, readable with kubectl logs. It also "
        "exposes three Prometheus metrics: api_requests_total (by endpoint, method and "
        "status), api_request_latency_seconds (histogram) and model_predictions_total (by "
        "outcome). Prometheus discovers both pods through the headless service and scrapes "
        "them every 15 seconds. Grafana loads a provisioned dashboard with six panels: pods "
        "up, predictions served, requests per second by endpoint, p95 prediction latency, "
        "responses by status code, and predictions by outcome.")
    report.p(
        "A traffic generator (scripts/send_traffic.py) sent 300 requests, one in ten "
        "deliberately invalid. The dashboard shows the resulting spike, the HTTP 200 and 422 "
        "responses, and 272 predictions served (the 270 valid generated requests plus a "
        "few manual test calls).")
    report.main_figures("monitoring")

    # ---- 12. Challenges
    report.h1("12. Challenges and lessons learned")
    report.bullets([
        "Ephemeral lab VMs: system installs, the Docker image and the Minikube cluster are "
        "lost when a lab VM is reset, while ~/workspace persists. Two idempotent scripts "
        "(setup_vm.sh and deploy_local.sh) now rebuild the toolchain and the deployment.",
        "MLflow 3: the file-based tracking store is in maintenance mode, so tracking moved "
        "to SQLite; the default model serialization refused a NumPy dtype, so the model is "
        "logged with cloudpickle.",
        "Reproducibility: kubectl was pinned to the cluster's Kubernetes version and every "
        "Python dependency to an exact version, because pickled models are version-sensitive.",
        "Authentication: GitHub no longer accepts passwords for git over HTTPS; the GitHub CLI "
        "browser login was used, with a token scoped to workflows for pushing CI files.",
        "Preventing leakage: imputation, scaling and encoding were placed inside the "
        "cross-validated pipeline rather than applied to the whole dataset beforehand.",
    ])

    # ---- 13. Limitations
    report.h1("13. Limitations and future work")
    report.bullets([
        "Small, single-cohort dataset (303 patients); results may not generalise and the "
        "61-patient test split gives wide confidence intervals.",
        "The decision threshold is fixed at 0.5; a threshold chosen for a target recall, and "
        "probability calibration, would suit a screening use case better.",
        "Training code has no unit tests of its own; a small smoke test would raise coverage.",
        "Monitoring covers service health only; data-drift and prediction-drift monitoring, a "
        "model registry, autoscaling, an Ingress with TLS and a Helm chart are natural next "
        "steps.",
    ])

    # ---- 14. Conclusion
    report.h1("14. Conclusion")
    report.p(
        "The project delivers the full lifecycle: a reproducible training pipeline with "
        "tracked experiments, a tested and automated CI pipeline, a containerised API "
        "deployed on Kubernetes, and monitoring with Prometheus and Grafana. The tuned "
        "Logistic Regression reached a cross-validated ROC-AUC of about 0.91 and a test "
        "ROC-AUC of 0.96, matching a Random Forest while remaining simple and interpretable.")
    report.p(f"Code repository: {REPO_URL}")

    report.h1("References")
    report.bullets([
        "UCI Machine Learning Repository, Heart Disease dataset: archive.ics.uci.edu",
        "scikit-learn, MLflow, FastAPI, Docker, Kubernetes/Minikube, Prometheus and Grafana "
        "official documentation.",
    ])

    report.appendix()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"saved {OUT}")
    print(f"figures: {report.fig_no}")
    if report.missing:
        print("missing images:", ", ".join(sorted(set(report.missing))))
    print("convert to PDF: soffice --headless --convert-to pdf --outdir report " + str(OUT))


if __name__ == "__main__":
    build()
